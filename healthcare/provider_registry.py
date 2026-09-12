"""healthcare/provider_registry.py — ADH service-provider registry import/export.

The ADH team maintains a working spreadsheet of the provider network. This
module reads that EXACT sheet into the `ServiceProvider` master and builds the
filtered Excel reports back out.

Two hard rules, from the brief + the 3-AI review (never clobber master data):
  1. Match ONLY on the AFA practice number, read as TEXT (leading zeros kept).
     Never match on name — two practices can share a name.
  2. Import is a PREVIEW then a COMMIT. `preview_import()` classifies every row
     as new / changed / unchanged / unmatched-blank and returns it for a human
     to approve; nothing is written until `commit_import()` is called with the
     rows the human approved.

No hard deletes anywhere — providers absent from a later sheet are left as-is
(deactivate is a deliberate, separate action), never deleted.
"""
from __future__ import annotations

import datetime as _dt
import io
from typing import Any

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import ServiceProvider, ServiceProviderApplication

# --- Exact spreadsheet header -> model field. Header text is matched
#     case-insensitively and whitespace-normalised so a stray trailing space
#     ("Contract Status ") or a re-cased header still lands. Invent nothing:
#     every key here is a real column from ADH's own sheet.
COLUMN_MAP: dict[str, str] = {
    "adh acceptance (ready)": "adh_acceptance",
    "practice": "practice_number",
    "discipline": "discipline",
    "prac name": "name",
    "email addr": "email",
    "town": "town",
    "contact no": "contact_number",
    "location": "location",
    "contract status": "contract_status",
    "welcome pack provided": "welcome_pack",
    "adh 'accepted here' sticker displayed": "sticker_displayed",
    "provider orientation": "provider_orientation",
    "provider onboarding link": "onboarding_link",
    "date contancted": "date_contacted",   # sheet's spelling, kept verbatim
    "comment": "comment",
    "vendor": "vendor_system",
}

# The importable text fields (everything except the derived readiness, the
# explicit QC gate, and lifecycle/provenance which are set by staff or code).
_IMPORT_FIELDS = [
    "adh_acceptance", "discipline", "name", "email", "town", "contact_number",
    "location", "contract_status", "welcome_pack", "sticker_displayed",
    "provider_orientation", "onboarding_link", "date_contacted", "comment",
    "vendor_system",
]


def _norm_header(v: Any) -> str:
    return " ".join(str(v or "").strip().lower().split())


def _cell(v: Any) -> str:
    """Render a cell as clean text. Numbers (e.g. a practice number Excel stored
    as a float) become plain ints, never '60178.0'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat()[:10]
    return str(v).strip()


def normalise_practice_number(v: Any) -> str:
    """Practice number as a stable TEXT key. Excel may deliver it as a float
    (60178.0) or a string with spaces; strip to the bare token but keep any
    genuine leading zeros a string carried."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_sheet(file_obj) -> tuple[list[dict], list[str]]:
    """Read the workbook's first sheet into a list of {field: value} dicts keyed
    by the COLUMN_MAP. Returns (rows, warnings). Rows with a blank practice
    number are dropped and counted in warnings — they cannot be matched."""
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return [], ["The sheet is empty."]

    # Map each column index to a model field via the header text.
    idx_to_field: dict[int, str] = {}
    for i, h in enumerate(header):
        field = COLUMN_MAP.get(_norm_header(h))
        if field:
            idx_to_field[i] = field

    warnings: list[str] = []
    missing = set(COLUMN_MAP.values()) - set(idx_to_field.values())
    if "practice_number" in missing:
        return [], ["No 'PRACTICE' column found — cannot import without the practice number key."]

    out: list[dict] = []
    blank_key = 0
    for r in rows_iter:
        rec: dict[str, str] = {}
        for i, field in idx_to_field.items():
            val = r[i] if i < len(r) else None
            if field == "practice_number":
                rec[field] = normalise_practice_number(val)
            else:
                rec[field] = _cell(val)
        # skip fully-empty rows
        if not any((rec.get(f) or "") for f in COLUMN_MAP.values()):
            continue
        if not rec.get("practice_number"):
            blank_key += 1
            continue
        out.append(rec)

    if blank_key:
        warnings.append(f"{blank_key} row(s) had no practice number and were skipped.")
    return out, warnings


def preview_import(file_obj) -> dict:
    """Classify every sheet row against the current registry WITHOUT writing.

    Returns {new, changed, unchanged, warnings, duplicates} where new/changed
    carry the field-level detail a human approves before commit.
    """
    rows, warnings = parse_sheet(file_obj)
    existing = {p.practice_number: p for p in ServiceProvider.objects.all()}

    seen: dict[str, int] = {}
    for row in rows:
        seen[row["practice_number"]] = seen.get(row["practice_number"], 0) + 1
    duplicates = [k for k, n in seen.items() if n > 1]
    if duplicates:
        warnings.append(
            f"{len(duplicates)} practice number(s) appear more than once in the "
            f"sheet: {', '.join(duplicates[:8])}{'…' if len(duplicates) > 8 else ''}. "
            "The last occurrence wins."
        )

    new_rows, changed_rows, unchanged = [], [], 0
    for row in rows:
        pn = row["practice_number"]
        cur = existing.get(pn)
        if cur is None:
            new_rows.append({"practice_number": pn, "name": row.get("name", ""), "values": row})
            continue
        diffs = {}
        for f in _IMPORT_FIELDS:
            old = getattr(cur, f, "") or ""
            newv = row.get(f, "") or ""
            if old != newv:
                diffs[f] = {"old": old, "new": newv}
        if diffs:
            changed_rows.append({
                "practice_number": pn, "name": cur.name, "changes": diffs, "values": row,
            })
        else:
            unchanged += 1

    return {
        "total_rows": len(rows),
        "new": new_rows,
        "changed": changed_rows,
        "unchanged": unchanged,
        "duplicates": duplicates,
        "warnings": warnings,
    }


def commit_import(rows: list[dict], *, source_file: str = "", user=None) -> dict:
    """Write approved rows into the registry. `rows` is a list of the per-row
    `values` dicts from preview_import (new + changed the human accepted).

    Existing rows are updated field-by-field; missing rows are created. The
    manual `adh_acceptance` column is written through, but derived readiness is
    computed live (properties) so it is always consistent. AuditableMixin writes
    the field-level audit row on each save.
    """
    now = timezone.now()
    created, updated = 0, 0
    # All-or-nothing: a mid-row DB error (e.g. an over-length value) must not
    # leave a half-written import behind. Roll the whole batch back instead.
    with transaction.atomic():
        for row in rows:
            pn = normalise_practice_number(row.get("practice_number"))
            if not pn:
                continue
            obj = ServiceProvider.objects.filter(practice_number=pn).first()
            is_new = obj is None
            if is_new:
                obj = ServiceProvider(practice_number=pn)
            for f in _IMPORT_FIELDS:
                if f in row:
                    setattr(obj, f, row.get(f, "") or "")
            # Migration seed (CREATE only): the team's existing "ADH Acceptance
            # (Ready) = YES" IS their prior QC sign-off, so carry it in as the
            # first qc_confirmed on a brand-new record — otherwise day-one
            # readiness reads as 0 and every prior YES shows as a mismatch. On a
            # RE-import of an existing record we never touch qc_confirmed, so an
            # in-system QC decision is authoritative and never clobbered.
            if is_new and (row.get("adh_acceptance", "") or "").strip().upper() == "YES":
                obj.qc_confirmed = True
            obj.source_file = source_file or obj.source_file
            obj.last_imported_at = now
            obj.save(audit_user=user, audit_description=(
                f"Imported {'new' if is_new else 'update'} from {source_file or 'spreadsheet'}"))
            created += int(is_new)
            updated += int(not is_new)
    return {"created": created, "updated": updated}


# ---------------------------------------------------------------------------
# Exports — the four filtered reports the brief asks for.
# ---------------------------------------------------------------------------

_EXPORT_HEADERS = [
    ("practice_number", "Practice"),
    ("name", "Prac Name"),
    ("discipline", "Discipline"),
    ("town", "Town"),
    ("email", "Email Addr"),
    ("contact_number", "Contact No"),
    ("location", "Location"),
    ("contract_status", "Contract Status"),
    ("afa_registered", "AFA Registered (derived)"),
    ("welcome_pack", "Welcome Pack Provided"),
    ("sticker_displayed", "'Accepted Here' Sticker"),
    ("provider_orientation", "Provider Orientation"),
    ("qc_confirmed", "QC Confirmed"),
    ("adh_ready", "ADH Ready (derived)"),
    ("adh_acceptance", "ADH Acceptance (manual)"),
    ("ready_mismatch", "Ready mismatch?"),
    ("vendor_system", "Vendor"),
    ("comment", "Comment"),
]

FILTERS = {
    "all": lambda qs: qs,
    "afa_registered": lambda qs: [p for p in qs if p.afa_registered == "Yes"],
    "adh_ready": lambda qs: [p for p in qs if p.adh_ready],
    "registered_not_qc": lambda qs: [p for p in qs if p.afa_registered == "Yes" and not p.qc_confirmed_flag],
}

FILTER_LABELS = {
    "all": "All providers",
    "afa_registered": "AFA-registered",
    "adh_ready": "ADH-ready",
    "registered_not_qc": "Registered, not yet QC'd",
}


def _display(obj, field: str) -> Any:
    val = getattr(obj, field)
    if isinstance(val, bool):
        return "Yes" if val else "No"
    return val if val is not None else ""


def build_export(filter_key: str = "all") -> tuple[bytes, str]:
    """Build an .xlsx of the chosen filtered report. Returns (bytes, filename)."""
    if filter_key not in FILTERS:
        filter_key = "all"
    providers = ServiceProvider.objects.filter(is_active=True).order_by("name")
    rows = FILTERS[filter_key](providers)

    wb = Workbook()
    ws = wb.active
    ws.title = FILTER_LABELS[filter_key][:31]
    ws.append([label for _, label in _EXPORT_HEADERS])
    for obj in rows:
        ws.append([_display(obj, f) for f, _ in _EXPORT_HEADERS])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = timezone.now().strftime("%Y-%m-%d")
    fname = f"ADH_service_providers_{filter_key}_{stamp}.xlsx"
    return buf.getvalue(), fname


def dashboard_counts() -> dict:
    """Readiness summary for the dashboard tiles — computed in Python because
    readiness is derived (not a DB column)."""
    providers = list(ServiceProvider.objects.filter(is_active=True))
    by_discipline: dict[str, int] = {}
    for p in providers:
        d = (p.discipline or "—").upper()
        by_discipline[d] = by_discipline.get(d, 0) + 1
    return {
        "total": len(providers),
        "afa_registered": sum(1 for p in providers if p.afa_registered == "Yes"),
        "afa_pending": sum(1 for p in providers if p.afa_registered == "Pending"),
        "adh_ready": sum(1 for p in providers if p.adh_ready),
        "registered_not_qc": sum(1 for p in providers if p.afa_registered == "Yes" and not p.qc_confirmed_flag),
        "mismatches": sum(1 for p in providers if p.ready_mismatch),
        "pending_applications": ServiceProviderApplication.objects.filter(status="pending").count(),
        "by_discipline": dict(sorted(by_discipline.items(), key=lambda kv: -kv[1])),
    }
