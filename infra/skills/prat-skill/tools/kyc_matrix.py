#!/usr/bin/env python3
"""KYC Collection Matrix — one-command pull from the Graphite V2 database.

Give it a list of policy numbers (xlsx / csv / txt) and it produces a branded
Excel showing, per policy: which KYC documents are on file, data-protection
consent, cancelled document, the agent, KYC compliance, store, and the most
recent document uploaded with its date. Source of truth = the Graphite DB
(customer_kyc + policies + users + stores + the document tables), reached via
ECS exec into the prod backend container. No browser, no rate limits — seconds.

PREREQUISITES (one-time, already set up 2026-06-25):
  * AWS Session Manager plugin installed (winget install Amazon.SessionManagerPlugin)
  * AWS profile `claude-cli` with ecs:ExecuteCommand on the graphite cluster
    (inline policy `claude-ecs-exec`).

USAGE:
  python kyc_matrix.py                              # input=Desktop\\KYC.xlsx, output auto-named
  python kyc_matrix.py --input "C:\\path\\list.xlsx" --output "C:\\path\\out.xlsx"
  python kyc_matrix.py --input list.txt            # one policy number per line

Reads policy numbers from the FIRST column (xlsx/csv) or one-per-line (txt);
a header cell like "Policy Number" is skipped automatically.
"""
import argparse, base64, json, os, re, subprocess, sys, datetime

PLUGIN_DIR = r"C:\Program Files\Amazon\SessionManagerPlugin\bin"
DEFAULTS = dict(profile="claude-cli", region="af-south-1", cluster="graphite-cluster",
                service="graphite-prod-backend", container="graphite-backend")
DESKTOP = os.path.join(os.path.expanduser("~"), "OneDrive - Alpha Direct Insurance", "Desktop")

PHP_TEMPLATE = r"""
error_reporting(0);
$nums = [__NUMS__];
$y = function($v){ return ($v !== null && $v !== '') ? 'Y' : ''; };
$pols = DB::table('policies as p')
  ->leftJoin('users as u','u.id','=','p.agent_id')
  ->leftJoin('stores as s','s.id','=','p.storeID')
  ->leftJoin('customer_kyc as k','k.customer_id','=','p.customer_id')
  ->whereIn('p.policyNumber',$nums)
  ->get(['p.id','p.policyNumber','u.firstName','u.lastName','s.name as store','k.compliance','k.created_at as kyc_created','k.omang','k.omangBack','k.passport','k.passport_back','k.passport_front_image','k.driving_license','k.driving_license_back','k.drivers_license_front_image','k.proof_residence','k.proof_of_address_image','k.proof_income','k.bank_statement_file_path','k.debit_authorization_form','k.data_protection_consent','k.Canceled_document']);
$ids = $pols->pluck('id')->filter()->all();
$last = [];
if (count($ids)) {
  $pa = DB::table('policy_attachments')->whereIn('policy_id',$ids)->selectRaw("policy_id, name as dname, created_at");
  $pd = DB::table('policy_documents')->whereIn('policy_id',$ids)->selectRaw("policy_id, 'Policy Schedule' as dname, created_at");
  $pk = DB::table('policy_kyc_documents')->whereIn('policy_id',$ids)->selectRaw("policy_id, doc_type as dname, created_at");
  $union = $pa->unionAll($pd)->unionAll($pk);
  $docs = DB::query()->fromSub($union,'d')->orderByDesc('created_at')->get();
  foreach ($docs as $d) { if (!isset($last[$d->policy_id])) $last[$d->policy_id] = ['name'=>$d->dname,'date'=>$d->created_at]; }
}
$lab = function($c){ if ((string)$c==='1') return 'KYC Compliant'; if ((string)$c==='2') return 'KYC Non Compliant'; return ''; };
$out = [];
$cands = [['omang','Omang ID Front'],['omangBack','Omang ID Back'],['passport','Passport'],['passport_back','Passport'],['passport_front_image','Passport'],['driving_license','Driving Licence'],['driving_license_back','Driving Licence'],['drivers_license_front_image','Driving Licence'],['proof_residence','Proof of Residence'],['proof_of_address_image','Proof of Residence'],['proof_income','Proof of Income'],['bank_statement_file_path','Bank Statement'],['debit_authorization_form','Debit Authorisation Form'],['data_protection_consent','Data Protection Consent']];
foreach ($pols as $r) {
  $ld = $last[$r->id] ?? null;
  if (!$ld) {
    foreach ($cands as $cc) { if (!empty($r->{$cc[0]})) { $ld = ['name'=>$cc[1], 'date'=>$r->kyc_created]; break; } }
  }
  $ld = $ld ?: ['name'=>'','date'=>''];
  $out[] = [
    'pn'=>$r->policyNumber,
    'oF'=>$y($r->omang), 'oB'=>$y($r->omangBack),
    'pp'=>$y($r->passport ?: ($r->passport_back ?: $r->passport_front_image)),
    'dl'=>$y($r->driving_license ?: ($r->driving_license_back ?: $r->drivers_license_front_image)),
    'pr'=>$y($r->proof_residence ?: $r->proof_of_address_image),
    'pi'=>$y($r->proof_income), 'bs'=>$y($r->bank_statement_file_path), 'da'=>$y($r->debit_authorization_form),
    'dpc'=>$y($r->data_protection_consent), 'canc'=>$y($r->Canceled_document),
    'agent'=>trim(($r->firstName ?? '').' '.($r->lastName ?? '')),
    'comp'=>$lab($r->compliance), 'store'=>$r->store ?? '',
    'lastdoc'=>$ld['name'], 'lastdate'=>$ld['date'],
  ];
}
echo "##"."ZZ"."##".base64_encode(json_encode($out))."##"."YY"."##";
"""

HEADER = ['Policy Number','Omang ID Front','Omang ID Back','Passport','Driving Licence',
          'Proof of Residence','Proof of Income','Bank Statement','Debit Authorisation Form',
          'Data Protection Consent','Cancelled Document','Agent','KYC Compliance','Store',
          'Last document attached','Last action date']
KEYS = ['oF','oB','pp','dl','pr','pi','bs','da','dpc','canc']  # the centered Y/blank doc cols


def read_numbers(path):
    ext = os.path.splitext(path)[1].lower()
    nums = []
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        ws = openpyxl.load_workbook(path, read_only=True).active
        for row in ws.iter_rows(min_row=1, max_col=1, values_only=True):
            if row[0] is not None:
                nums.append(str(row[0]).strip())
    elif ext == ".csv":
        import csv
        with open(path, encoding="utf-8-sig", newline="") as f:
            for row in csv.reader(f):
                if row and row[0].strip():
                    nums.append(row[0].strip())
    else:
        with open(path, encoding="utf-8-sig") as f:
            nums = [ln.strip() for ln in f if ln.strip()]
    # drop a header-ish first cell
    if nums and nums[0].lower().replace(" ", "").replace("_", "") in ("policynumber", "policy", "policyno"):
        nums = nums[1:]
    return [n for n in nums if n]


def run_query(nums, args):
    env = os.environ.copy()
    env["PATH"] = PLUGIN_DIR + os.pathsep + env.get("PATH", "")
    php = PHP_TEMPLATE.replace("__NUMS__", ",".join('"%s"' % n.replace('"', '') for n in nums))
    b64 = base64.b64encode(php.encode()).decode()
    task = subprocess.check_output(
        ["aws","ecs","list-tasks","--cluster",args.cluster,"--service-name",args.service,
         "--profile",args.profile,"--region",args.region,"--query","taskArns[0]","--output","text"],
        env=env, text=True, encoding="utf-8", errors="replace").strip()
    if not task or task == "None":
        sys.exit("ERROR: no running task found for service " + args.service)
    cmd = ["aws","ecs","execute-command","--cluster",args.cluster,"--task",task,
           "--container",args.container,"--interactive",
           "--command", "sh -c 'echo %s | base64 -d | php artisan tinker'" % b64,
           "--profile",args.profile,"--region",args.region]
    raw = subprocess.run(cmd, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace").stdout
    m = re.search(r"##ZZ##(.*?)##YY##", raw, re.S)
    if not m:
        sys.exit("ERROR: no data returned from DB. Raw tail:\n" + raw[-600:])
    clean = re.sub(r"[^A-Za-z0-9+/=]", "", m.group(1))
    return json.loads(base64.b64decode(clean).decode())


def fmtdate(s):
    s = (s or "").strip()
    if not s:
        return ""
    try:
        return datetime.datetime.strptime(s, "%Y-%m-%d %H:%M:%S").strftime("%d %b %Y %H:%M")
    except Exception:
        return s


def _tc(s):
    """Title-case agent names (DB stores some uppercase); keep it simple."""
    return " ".join(w.capitalize() for w in (s or "").split())


def build_excel(nums, by_pn, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    NAVY="0D1B2A"; ORANGE="F4A623"; WHITE="FFFFFF"; LIGHT="F4F6F8"; GREEN="1E7A34"; AMBER="B36B00"
    wb = Workbook(); ws = wb.active; ws.title = "KYC Matrix"
    thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
    n = len(HEADER); date_col = n
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    t = ws.cell(1, 1, "Alpha Direct  -  KYC Collection Matrix   .   %d policies   .   source: Graphite V2 database, %s"
                % (len(nums), datetime.date.today().strftime("%d %b %Y")))
    t.font = Font(name="Book Antiqua", size=13, bold=True, color=WHITE); t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1); ws.row_dimensions[1].height = 28
    for c, h in enumerate(HEADER, start=1):
        cell = ws.cell(2, c, h); cell.font = Font(name="Book Antiqua", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color=ORANGE))
    ws.row_dimensions[2].height = 40
    doc_cols = set(range(2, 12))
    blank = {"pn":"","oF":"","oB":"","pp":"","dl":"","pr":"","pi":"","bs":"","da":"","dpc":"","canc":"","agent":"","comp":"","store":"","lastdoc":"","lastdate":""}
    for ri, pn in enumerate(nums, start=3):
        d = by_pn.get(pn, blank)
        vals = [pn, d["oF"],d["oB"],d["pp"],d["dl"],d["pr"],d["pi"],d["bs"],d["da"],d["dpc"],d["canc"],
                _tc(d["agent"]), d["comp"], d["store"], d["lastdoc"], fmtdate(d["lastdate"])]
        stripe = (ri % 2 == 1)
        for ci in range(1, n+1):
            val = vals[ci-1]
            cell = ws.cell(ri, ci, val); cell.border = border
            if stripe: cell.fill = PatternFill("solid", fgColor=LIGHT)
            if ci in doc_cols:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Calibri", size=10, bold=True, color=GREEN) if val == "Y" else Font(name="Calibri", size=10, color="BBBBBB")
            elif ci == date_col:
                cell.alignment = Alignment(horizontal="center", vertical="center"); cell.font = Font(name="Calibri", size=10)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1); cell.font = Font(name="Calibri", size=10)
    widths = [16,9,9,8,9,12,11,11,13,13,13,20,15,24,26,17]
    for i, w in enumerate(widths[:n], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B3"; ws.auto_filter.ref = "A2:%s%d" % (get_column_letter(n), 2+len(nums))
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="Pull KYC matrix from Graphite DB into Excel.")
    ap.add_argument("--input", default=os.path.join(DESKTOP, "KYC.xlsx"))
    ap.add_argument("--output", default=None)
    for k, v in DEFAULTS.items():
        ap.add_argument("--" + k, default=v)
    args = ap.parse_args()
    if not os.path.exists(args.input):
        sys.exit("ERROR: input not found: " + args.input)
    out = args.output or os.path.join(DESKTOP, "KYC_matrix_%s.xlsx" % datetime.date.today().strftime("%Y%m%d"))

    nums = read_numbers(args.input)
    print("policies in:", len(nums))
    data = run_query(nums, args)
    by_pn = {r["pn"]: r for r in data}
    print("policies returned from DB:", len(by_pn))
    missing = [n for n in nums if n not in by_pn]
    if missing:
        print("WARNING: %d not found in DB (left blank):" % len(missing), missing[:10])
    build_excel(nums, by_pn, out)
    have = sum(1 for n in nums if n in by_pn)
    print("WROTE %s  (%d rows, %d matched)" % (out, len(nums), have))


if __name__ == "__main__":
    main()
