import openpyxl, json, re
from collections import defaultdict
TMP = r"C:\Users\PrathapAsus\OneDrive - Alpha Direct Insurance\Gods Eye\_shay_tmp.xlsx"
STRUCT = r"C:\Users\PrathapAsus\OneDrive - Alpha Direct Insurance\Gods Eye\_shay_struct.json"
OUT = r"C:\Users\PrathapAsus\OneDrive - Alpha Direct Insurance\Desktop\Shaysons_Graphite_Import.xlsx"

st = json.load(open(STRUCT, encoding="utf-8"))
policy_locs = sorted(set(r['loc'] for r in st['pcd']))
def nloc(s):
    s=str(s).upper(); s=re.sub(r'\b(PLOT|NO|WARD|PROPRIETARY|PTY|LTD)\b','',s); return re.sub(r'[^A-Z0-9]','',s)
pn={nloc(l):l for l in policy_locs}; OVR={'SHAYSONSHO':'SHAYSONS HQ'}
def maploc(x):
    base=re.sub(r'\s*2026\s*$','',str(x)).strip(); n=nloc(base)
    if n in OVR: return OVR[n]
    if n in pn: return pn[n]
    for k,v in pn.items():
        if k and (k in n or n in k): return v
    return None

SECTION = {
 'fire & perils':'FIRE','fire and perils':'FIRE','fire':'FIRE',
 'business interruption':'BUSINESSINTERUPTION',
 'electronic equipment':'ELECTRONICEQUIPMENT',
 'theft':'THEFT','fidelity guarantee':'FIDELITYGUARANTEE','money':'MONEY',
 'accidental damage':'ACCIDENTALDAMAGE','office contents':'OFFICECONTENTS',
 'workmans compensation':'WORKERSCOMPENSATION','workers compensation':'WORKERSCOMPENSATION',
 'public liability':'PUBLICLIABILITY','goods in transit':'GOODSINTRANSIT',
 'commercial motor':'COMMERCIALMOTOR','motor':'COMMERCIALMOTOR','motor fleet':'COMMERCIALMOTOR','motor comprehensive':'COMMERCIALMOTOR',
}
KNOWN_NONSECTION = {'insured:','main physical address:','postal address:','period:','cover','changes made',
 'indemnity period - 12 months','description','sum insured'}

wb = openpyxl.load_workbook(TMP, data_only=True)
rows = defaultdict(list)         # cov -> list of (loc,label,si,q)
unknown_headers = set()
ramok = 0; ramok_sum = 0.0; new_locs={}
for ws in wb.worksheets:
    if not ws.title.endswith('2026'): continue
    loc = maploc(ws.title); cur=None
    if loc is None:
        loc = re.sub(r'\s*2026\s*$','',ws.title).strip().upper(); new_locs.setdefault(loc,0.0)
    for r in ws.iter_rows(min_row=5, values_only=True):
        a=r[0]; b=r[1] if len(r)>1 else None; e=r[4] if len(r)>4 else None
        if a is None: continue
        al=re.sub(r'\s+',' ',str(a).strip().lower())
        if al in SECTION: cur=SECTION[al]; continue
        if 'total' in al: continue
        if isinstance(e,(int,float)) and e!=0:
            si=float(b) if isinstance(b,(int,float)) else None
            if cur is None:
                unknown_headers.add('(no section)'); continue
            rows[cur].append((loc,str(a).strip(),si,round(float(e),2)))
            if loc in new_locs: new_locs[loc]+=round(float(e),2)
        elif not isinstance(b,(int,float)) and not isinstance(e,(int,float)) and al not in KNOWN_NONSECTION and len(al)<45 and al:
            unknown_headers.add(str(a).strip())

# write import workbook (one sheet per coverage code)
out = openpyxl.Workbook(); out.remove(out.active)
PLATE = re.compile(r'\bB\s?\d{2,3}\s?[A-Z]{2,3}\b')
# action 43815's existing 42 fleet plates (trimmed) — lines matching these reuse the vehicle
EXIST = set("""B160AVR B678AUJ B112ARY B652ATZ B865ASY B900AVB B979AVH B757AWS B668AWS B665AWS
B721AZJ B692AZJ B945AZD B203BAX B513BCH B233BBE B696BFX B641BFX B673BFX B761BLC B776BLC B722BLD
B720BLD B712BLD B927BGI B924BGI B649BMO B301BIS B520BKS B516BKS B921BUK B950BUM B952BUM S14846
S35699 B243BUP B166BUI B724AZJ""".split())
EXIST |= {"N/A 1","N/A 2","TBA 1","TBA 3"}
new_vehicles=[]; fk=[0]
def mkreg(label):
    U=label.upper(); m=PLATE.search(U); plate=(m.group(0).replace(' ','') if m else None)
    if 'MOBILE TRAILER' in U: return 'TBA 1',None
    if plate=='B6737BFX' or 'B6737BFX' in U.replace(' ',''): return 'B673BFX',None
    if plate and plate in EXIST: return plate,None
    if plate: return plate,{'reg':plate,'make':label.strip()[:45],'model':''}
    fk[0]+=1; r=f"FORKLIFT {fk[0]}"; return r,{'reg':r,'make':label.strip()[:45],'model':''}
# policy's existing primary sub-coverage name per coverage (exact strings incl. typos)
PRIMARY = {
 'FIRE':'Plant and Machinery',
 'BUSINESSINTERUPTION':'Gross Profit',
 'MONEY':'Major Limit',
 'ACCIDENTALDAMAGE':'All the Insured property as defined in this section(First Loss)',
 'OFFICECONTENTS':'Contents',
 'THEFT':'Damage to Buildings',
 'WORKERSCOMPENSATION':'Employers Liablity (Common Law Liability)',
 'ELECTRONICEQUIPMENT':'List of items',
 'PUBLICLIABILITY':'Full address of Premises Covered',
 'GOODSINTRANSIT':'Estimated Annual Carry',
}
grand=0.0; summary=[]
def agg_by_loc(lines):
    a={}
    for loc,label,si,q in lines:
        x=a.setdefault(loc,[0.0,0.0]); x[0]+=(si or 0); x[1]+=q
    return a
for cov, lines in rows.items():
    ws = out.create_sheet(cov[:31])
    if cov=='COMMERCIALMOTOR':
        ws.append(['Risk Address','Registration No','Coverage Value','Calculated Value'])
        for loc,label,si,q in lines:
            reg,nv=mkreg(label)
            if nv: nv['si']=si; new_vehicles.append(nv)
            ws.append([loc,reg,si,q]); grand+=q
        n=len(lines)
    elif cov=='FIDELITYGUARANTEE':
        ws.append(['Risk Address','cover_type','amount_to_be_guaranteed','Premium'])
        a=agg_by_loc(lines)
        for loc,(si,q) in a.items():
            ws.append([loc,'Fidelity',round(si,2),round(q,2)]); grand+=round(q,2)
        n=len(a)
    else:
        name=PRIMARY.get(cov,'Sum Insured')
        ws.append(['Risk Address','Coverage','Limit','Premium'])
        a=agg_by_loc(lines)
        for loc,(si,q) in a.items():
            ws.append([loc,name,round(si,2),round(q,2)]); grand+=round(q,2)
        n=len(a)
    t=sum(l[3] for l in lines); summary.append((cov,n,round(t,2)))
out.save(OUT)
json.dump(new_vehicles, open(r"C:\Users\PrathapAsus\OneDrive - Alpha Direct Insurance\Gods Eye\_newveh.json","w"), indent=1)

print("WROTE", OUT)
print("NEW VEHICLES to create:", len(new_vehicles))
for nv in new_vehicles: print(f"   {nv['reg']:14} SI={nv['si']}  {nv['make'][:34]}")
for cov,n,t in sorted(summary):
    print(f"  {cov:22} lines={n:4}  premium={t:>12,.2f}")
json.dump(new_locs, open(r"C:\Users\PrathapAsus\OneDrive - Alpha Direct Insurance\Gods Eye\_newlocs.json","w"), indent=1)
print(f"  GRAND TOTAL (all locations) = {grand:,.2f}   (target 184,023.50)")
print(f"  NEW locations to create: {new_locs}")
if unknown_headers: print("  UNRECOGNISED section headers:", sorted(unknown_headers))
