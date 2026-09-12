#!/usr/bin/env python3
"""KYC Activity Log enrichment — adds the chronological KYC activity history to a KYC matrix Excel.

Companion to kyc_matrix.py. Reads the policy numbers from an existing KYC matrix file,
pulls the complete kyc_activity_log per policy from the Graphite V2 DB (via ECS exec),
and adds:
  - a "KYC Activity Log" sheet: every KYC event, chronological per policy
    (document approvals/rejections, compliance changes, overall decisions; with date,
    reviewer, status and reason) — the same Activity Log shown on the KYC review screen.
  - five summary columns on the main sheet: Latest KYC Activity / Latest Status /
    Latest By / Latest Activity Date / Activity Events (count).
Idempotent: re-running replaces the sheet and reuses the summary columns.
Source of truth = Graphite V2 `kyc_activity_log` (reviewer ids resolved to names via `users`).

USAGE (Windows laptop only):
  python kyc_activity_log.py                              # newest Desktop\\KYC_matrix_*.xlsx, in place
  python kyc_activity_log.py --input "C:\\path\\matrix.xlsx" [--output "C:\\path\\out.xlsx"]

PREREQS (one-time, set up 2026-06-25): AWS Session Manager plugin installed +
`claude-cli` profile with ecs:ExecuteCommand on the graphite cluster (policy claude-ecs-exec).
"""
import argparse, base64, json, os, re, subprocess, sys, datetime, glob
from collections import defaultdict

PLUGIN_DIR = r"C:\Program Files\Amazon\SessionManagerPlugin\bin"
DEFAULTS = dict(profile="claude-cli", region="af-south-1", cluster="graphite-cluster",
                service="graphite-prod-backend", container="graphite-backend")
DESKTOP = os.path.join(os.path.expanduser("~"), "OneDrive - Alpha Direct Insurance", "Desktop")
POLPAT = re.compile(r"^[A-Z]{2,6}\d{6,}$")

PHP_TEMPLATE = r"""
error_reporting(0);
$nums = [__NUMS__];
$pols = DB::table('policies')->whereIn('policyNumber',$nums)->get(['policyNumber','customer_id']);
$polByCust = [];
foreach($pols as $p){ if($p->customer_id) $polByCust[$p->customer_id][] = $p->policyNumber; }
$cids = array_keys($polByCust);
$acts = DB::table('kyc_activity_log')->whereIn('customer_id',$cids)->orderBy('action_performed_at')->orderBy('id')->get(['customer_id','action_performed_at','log_name','status','compliance','action_perfomed_by','description','reason']);
$uids=[]; foreach($acts as $a){ if(is_numeric($a->action_perfomed_by)) $uids[(int)$a->action_perfomed_by]=1; }
$uname=[]; if(count($uids)){ foreach(DB::table('users')->whereIn('id',array_keys($uids))->get(['id','firstName','lastName']) as $u){ $uname[$u->id]=trim(($u->firstName??'').' '.($u->lastName??'')); } }
$ev=[];
foreach($acts as $a){ $by = is_numeric($a->action_perfomed_by) ? ($uname[(int)$a->action_perfomed_by] ?? ('user#'.$a->action_perfomed_by)) : $a->action_perfomed_by; $ev[]=[$a->customer_id,$a->action_performed_at,$a->log_name,$a->description,$a->status,$a->compliance,$by,$a->reason]; }
echo "##"."ZZ"."##".base64_encode(json_encode(['map'=>$polByCust,'events'=>$ev]))."##"."YY"."##";
"""


def read_numbers(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True).active
    nums = []
    for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
        v = row[0]
        if v is not None and POLPAT.match(str(v).strip()):
            nums.append(str(v).strip())
    return nums


def run_query(nums, args):
    env = os.environ.copy()
    env["PATH"] = PLUGIN_DIR + os.pathsep + env.get("PATH", "")
    php = PHP_TEMPLATE.replace("__NUMS__", ",".join('"%s"' % n.replace('"', '') for n in nums))
    b64 = base64.b64encode(php.encode()).decode()
    task = subprocess.check_output(
        ["aws","ecs","list-tasks","--cluster",args.cluster,"--service-name",args.service,
         "--profile",args.profile,"--region",args.region,"--query","taskArns[0]","--output","text"],
        env=env, text=True, encoding="utf-8", errors="replace").strip()
    if not task or task == "None":
        sys.exit("ERROR: no running task for service " + args.service)
    cmd = ["aws","ecs","execute-command","--cluster",args.cluster,"--task",task,
           "--container",args.container,"--interactive",
           "--command", "sh -c 'echo %s | base64 -d | php artisan tinker'" % b64,
           "--profile",args.profile,"--region",args.region]
    raw = subprocess.run(cmd, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace").stdout
    m = re.search(r"##ZZ##(.*?)##YY##", raw, re.S)
    if not m:
        sys.exit("ERROR: no data returned from DB. Raw tail:\n" + raw[-600:])
    clean64 = re.sub(r"[^A-Za-z0-9+/=]", "", m.group(1))
    return json.loads(base64.b64decode(clean64).decode())


def clean(s):
    if s is None:
        return ""
    s = str(s).replace("<br>", " | ").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip()


def enrich(in_path, out_path, data):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    NAVY="0D1B2A"; ORANGE="F4A623"; WHITE="FFFFFF"; LIGHT="F4F6F8"; GREEN="1E7A34"; RED="B00020"
    thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin,right=thin,top=thin,bottom=thin)

    cust_to_pols = data["map"]
    ev_by_cust = defaultdict(list)
    for e in data["events"]:
        ev_by_cust[str(e[0])].append(e)
    pol_to_cust = {}
    for cust, pols in cust_to_pols.items():
        for pn in pols:
            pol_to_cust[pn] = str(cust)

    wb = openpyxl.load_workbook(in_path)
    ws = wb.worksheets[0]
    # locate header row (col A == "Policy Number")
    hdr_row = 2
    for r in range(1, min(6, ws.max_row)+1):
        if str(ws.cell(r,1).value).strip().lower() == "policy number":
            hdr_row = r; break

    # ---- summary columns (idempotent) ----
    sumhdr = ["Latest KYC Activity","Latest Status","Latest By","Latest Activity Date","Activity Events"]
    existing = {ws.cell(hdr_row,c).value: c for c in range(1, ws.max_column+1)}
    if all(h in existing for h in sumhdr):
        cols = [existing[h] for h in sumhdr]
    else:
        s = ws.max_column + 1
        cols = [s+i for i in range(len(sumhdr))]
        for i, h in enumerate(sumhdr):
            c = ws.cell(hdr_row, cols[i], h)
            c.font = Font(name="Book Antiqua", bold=True, color=WHITE, size=10)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(left=thin,right=thin,top=thin,bottom=Side(style="medium", color=ORANGE))
    for r in range(hdr_row+1, ws.max_row+1):
        pn = str(ws.cell(r,1).value).strip()
        if not POLPAT.match(pn):
            continue
        evs = ev_by_cust.get(pol_to_cust.get(pn), [])
        if evs:
            last = evs[-1]
            vals = [clean(last[3]), clean(last[4]), clean(last[6]), clean(last[1]), len(evs)]
        else:
            vals = ["(no KYC activity log)", "", "", "", 0]
        for i, v in enumerate(vals):
            cell = ws.cell(r, cols[i], v)
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            cell.font = (Font(name="Calibri", size=10, bold=True, color=RED)
                         if (i == 1 and isinstance(v, str) and "reject" in v.lower())
                         else Font(name="Calibri", size=10))
    widths = [26, 14, 18, 20, 11]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(cols[i])].width = w
    ws.auto_filter.ref = "A%d:%s%d" % (hdr_row, get_column_letter(ws.max_column), ws.max_row)

    # ---- activity log sheet (idempotent: replace) ----
    if "KYC Activity Log" in wb.sheetnames:
        del wb["KYC Activity Log"]
    al = wb.create_sheet("KYC Activity Log")
    H = ["Policy Number","Customer ID","Date / Time","Section","Activity","Status","Compliance","Performed By","Reason / Detail"]
    al.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(H))
    t = al.cell(1,1, "Alpha Direct  -  KYC Activity Log (chronological, per policy)   ·   source: Graphite V2 kyc_activity_log   ·   %s" % datetime.date.today().strftime("%d %b %Y"))
    t.font = Font(name="Book Antiqua", size=13, bold=True, color=WHITE); t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1); al.row_dimensions[1].height = 28
    for c, h in enumerate(H, start=1):
        cell = al.cell(2, c, h); cell.font = Font(name="Book Antiqua", bold=True, color=WHITE, size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin,right=thin,top=thin,bottom=Side(style="medium", color=ORANGE))
    al.row_dimensions[2].height = 30
    rownum = 3; total = 0
    for r in range(hdr_row+1, ws.max_row+1):
        pn = str(ws.cell(r,1).value).strip()
        if not POLPAT.match(pn):
            continue
        for e in ev_by_cust.get(pol_to_cust.get(pn), []):
            row = [pn, e[0], clean(e[1]), clean(e[2]), clean(e[3]), clean(e[4]), e[5], clean(e[6]), clean(e[7])]
            stripe = (rownum % 2 == 1)
            for ci, v in enumerate(row, start=1):
                cell = al.cell(rownum, ci, v); cell.border = border
                if stripe: cell.fill = PatternFill("solid", fgColor=LIGHT)
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=(ci in (5,9)), indent=1)
                f = Font(name="Calibri", size=9)
                if ci == 5 and isinstance(v, str) and "reject" in v.lower(): f = Font(name="Calibri", size=9, bold=True, color=RED)
                elif ci == 5 and isinstance(v, str) and "approv" in v.lower(): f = Font(name="Calibri", size=9, color=GREEN)
                cell.font = f
            rownum += 1; total += 1
    for i, w in enumerate([16,12,18,20,55,14,11,20,55], start=1):
        al.column_dimensions[get_column_letter(i)].width = w
    al.freeze_panes = "A3"
    al.auto_filter.ref = "A2:%s%d" % (get_column_letter(len(H)), max(2, rownum-1))
    wb.save(out_path)
    return total


def main():
    ap = argparse.ArgumentParser(description="Add KYC activity-log history to a KYC matrix Excel.")
    ap.add_argument("--input", default=None, help="KYC matrix xlsx (default: newest Desktop\\KYC_matrix_*.xlsx)")
    ap.add_argument("--output", default=None, help="output xlsx (default: in place)")
    for k, v in DEFAULTS.items():
        ap.add_argument("--" + k, default=v)
    args = ap.parse_args()
    inp = args.input
    if not inp:
        cands = glob.glob(os.path.join(DESKTOP, "KYC_matrix*.xlsx"))
        if not cands:
            sys.exit("ERROR: no input given and no Desktop\\KYC_matrix*.xlsx found")
        inp = max(cands, key=os.path.getmtime)
    if not os.path.exists(inp):
        sys.exit("ERROR: input not found: " + inp)
    out = args.output or inp
    print("input:", inp)
    nums = read_numbers(inp)
    print("policies:", len(nums))
    data = run_query(nums, args)
    print("activity events pulled:", len(data["events"]), "| customers:", len(data["map"]))
    total = enrich(inp, out, data)
    print("WROTE %s  (%d activity rows)" % (out, total))


if __name__ == "__main__":
    main()
